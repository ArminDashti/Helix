"""Import CSV, Excel, JSON, and SQLite files into a local SQLite analytics database."""

from __future__ import annotations

import csv
import json
import re
import sqlite3
import uuid
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, BinaryIO

from django.conf import settings

from .sample_database import sample_data_dir

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ROWS_PER_TABLE = 100_000

_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_SANITIZE = re.compile(r"[^A-Za-z0-9_]+")

CSV_EXTENSIONS = {".csv", ".tsv", ".txt"}
EXCEL_EXTENSIONS = {".xlsx", ".xls"}
JSON_EXTENSIONS = {".json", ".jsonl", ".ndjson"}
SQLITE_EXTENSIONS = {".sqlite", ".sqlite3", ".db"}


def imports_dir() -> Path:
    path = sample_data_dir() / "imports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _sanitize_table_name(raw: str, *, fallback: str = "data") -> str:
    base = Path(raw).stem if "." in raw else raw
    cleaned = _SANITIZE.sub("_", (base or "").strip()).strip("_")
    if not cleaned:
        cleaned = fallback
    if cleaned[0].isdigit():
        cleaned = f"t_{cleaned}"
    return cleaned[:63]


def _unique_table_name(base: str, used: set[str]) -> str:
    name = _sanitize_table_name(base)
    if name not in used and _IDENT_RE.match(name):
        used.add(name)
        return name
    idx = 2
    while True:
        candidate = f"{name}_{idx}"[:63]
        if candidate not in used:
            used.add(candidate)
            return candidate
        idx += 1


def _is_int(text: str) -> bool:
    try:
        int(text)
        return True
    except (TypeError, ValueError):
        return False


def _is_float(text: str) -> bool:
    try:
        float(text)
        return True
    except (TypeError, ValueError):
        return False


def _column_type(values: list[Any]) -> str:
    non_null = [
        v
        for v in values
        if v is not None and (not isinstance(v, str) or v.strip() != "")
    ]
    if not non_null:
        return "TEXT"
    as_str = [str(v).strip() for v in non_null]
    if all(_is_int(s) for s in as_str):
        return "INTEGER"
    if all(_is_float(s) for s in as_str):
        return "REAL"
    return "TEXT"


def _normalize_cell(value: Any, col_type: str) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and col_type == "INTEGER":
        if value.is_integer():
            return int(value)
    text = str(value).strip()
    if text == "":
        return None
    if col_type == "INTEGER" and _is_int(text):
        return int(text)
    if col_type == "REAL" and _is_float(text):
        return float(text)
    return text


def _write_table(
    conn: sqlite3.Connection,
    table: str,
    columns: list[str],
    rows: list[dict[str, Any]],
) -> None:
    if not columns:
        raise ValueError(f"Table {table} has no columns")
    if not rows:
        raise ValueError(f"Table {table} has no data rows")
    if len(rows) > MAX_ROWS_PER_TABLE:
        raise ValueError(
            f"Table {table} exceeds {MAX_ROWS_PER_TABLE:,} rows; split the file or filter rows."
        )
    col_types = {col: _column_type([row.get(col) for row in rows]) for col in columns}
    quoted_cols = ", ".join(f'"{c}"' for c in columns)
    decl = ", ".join(f'"{c}" {col_types[c]}' for c in columns)
    conn.execute(f'CREATE TABLE "{table}" ({decl})')
    placeholders = ", ".join("?" for _ in columns)
    insert_sql = f'INSERT INTO "{table}" ({quoted_cols}) VALUES ({placeholders})'
    payload = [
        tuple(_normalize_cell(row.get(col), col_types[col]) for col in columns)
        for row in rows
    ]
    conn.executemany(insert_sql, payload)


def _decode_text(upload: BinaryIO) -> str:
    raw = upload.read()
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit")
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode text file (try UTF-8 CSV)")


def _read_csv_table(upload: BinaryIO, table_name: str) -> tuple[str, list[str], list[dict[str, Any]]]:
    text = _decode_text(upload)
    if not text.strip():
        raise ValueError("CSV file is empty")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError("CSV has no header row")
    columns = [_sanitize_table_name(h, fallback=f"col_{i}") for i, h in enumerate(reader.fieldnames)]
    # ponytail: duplicate headers after sanitize → suffix
    seen: dict[str, int] = {}
    unique_cols: list[str] = []
    for col in columns:
        if col not in seen:
            seen[col] = 1
            unique_cols.append(col)
        else:
            seen[col] += 1
            unique_cols.append(f"{col}_{seen[col]}"[:63])
    rows: list[dict[str, Any]] = []
    for row in reader:
        if len(rows) >= MAX_ROWS_PER_TABLE:
            break
        item = {}
        for src, dest in zip(reader.fieldnames, unique_cols):
            item[dest] = row.get(src)
        if any(v is not None and str(v).strip() != "" for v in item.values()):
            rows.append(item)
    if not rows:
        raise ValueError("CSV has no data rows")
    return table_name, unique_cols, rows


def _read_json_tables(upload: BinaryIO) -> dict[str, list[dict[str, Any]]]:
    text = _decode_text(upload)
    if not text.strip():
        raise ValueError("JSON file is empty")
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError("Invalid JSON file") from exc
    tables: dict[str, list[dict[str, Any]]] = {}
    if isinstance(data, list):
        if not data or not all(isinstance(row, dict) for row in data):
            raise ValueError("JSON array must contain objects")
        tables["data"] = data[:MAX_ROWS_PER_TABLE]
        return tables
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, list) and value and all(isinstance(row, dict) for row in value):
                tables[str(key)] = value[:MAX_ROWS_PER_TABLE]
        if tables:
            return tables
    raise ValueError("JSON must be an array of objects or an object of arrays of objects")


def _read_jsonl_table(upload: BinaryIO, table_name: str) -> tuple[str, list[str], list[dict[str, Any]]]:
    text = _decode_text(upload)
    rows: list[dict[str, Any]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if len(rows) >= MAX_ROWS_PER_TABLE:
            break
        try:
            item = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError("Invalid JSON Lines row") from exc
        if not isinstance(item, dict):
            raise ValueError("JSON Lines rows must be objects")
        rows.append(item)
    if not rows:
        raise ValueError("JSON Lines file has no rows")
    columns = list(rows[0].keys())
    for row in rows[1:]:
        for key in row:
            if key not in columns:
                columns.append(key)
    safe_cols = [_sanitize_table_name(c, fallback=f"col_{i}") for i, c in enumerate(columns)]
    return table_name, safe_cols, [{safe_cols[i]: row.get(columns[i]) for i in range(len(columns))} for row in rows]


def _read_excel_tables(upload: BinaryIO) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel import requires openpyxl on the API host") from exc
    raw = upload.read()
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit")
    wb = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
    tables: dict[str, list[dict[str, Any]]] = {}
    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        headers = [
            _sanitize_table_name(str(h or f"col_{i}"), fallback=f"col_{i}")
            for i, h in enumerate(header_row or ())
        ]
        if not any(h for h in header_row or ()):
            continue
        data_rows: list[dict[str, Any]] = []
        for row in rows_iter:
            if len(data_rows) >= MAX_ROWS_PER_TABLE:
                break
            if row is None:
                continue
            item = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
            if any(v is not None and str(v).strip() != "" for v in item.values()):
                data_rows.append(item)
        if data_rows:
            tables[sheet.title or "Sheet1"] = data_rows
    wb.close()
    if not tables:
        raise ValueError("Excel workbook has no data")
    return tables


def _build_sqlite_from_tables(tables: dict[str, list[dict[str, Any]]], dest: Path) -> list[str]:
    used: set[str] = set()
    created: list[str] = []
    with sqlite3.connect(str(dest)) as conn:
        for raw_name, rows in tables.items():
            table = _unique_table_name(raw_name, used)
            columns = list(rows[0].keys()) if rows else []
            for row in rows[1:]:
                for key in row:
                    if key not in columns:
                        columns.append(key)
            safe_cols = [_sanitize_table_name(c, fallback=f"col_{i}") for i, c in enumerate(columns)]
            normalized = [{safe_cols[i]: row.get(columns[i]) for i in range(len(columns))} for row in rows]
            _write_table(conn, table, safe_cols, normalized)
            created.append(table)
        conn.commit()
    if not created:
        raise ValueError("No tables were created from the file")
    return created


def import_uploaded_file(upload: Any) -> tuple[Path, list[str]]:
    """Save upload to SQLite under data/imports/. Returns (sqlite path, table names)."""
    name = str(getattr(upload, "name", "") or "upload.dat")
    suffix = Path(name).suffix.lower()
    if suffix not in (
        CSV_EXTENSIONS | EXCEL_EXTENSIONS | JSON_EXTENSIONS | SQLITE_EXTENSIONS
    ):
        raise ValueError(
            "Unsupported file type. Use CSV, TSV, Excel (.xlsx), JSON, JSON Lines, or SQLite (.db)."
        )

    stem = Path(name).stem
    dest = imports_dir() / f"helix-import-{uuid.uuid4().hex[:12]}-{stem}.sqlite"
    if dest.exists():
        dest.unlink()

    if suffix in SQLITE_EXTENSIONS:
        raw = upload.read()
        if len(raw) > MAX_UPLOAD_BYTES:
            raise ValueError(f"File exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit")
        if hasattr(upload, "seek"):
            upload.seek(0)
        dest.write_bytes(raw)
        with sqlite3.connect(str(dest)) as conn:
            rows = conn.execute(
                """
                SELECT name FROM sqlite_master
                WHERE type='table' AND name NOT LIKE 'sqlite_%'
                ORDER BY name
                """
            ).fetchall()
        tables = [str(r[0]) for r in rows]
        if not tables:
            dest.unlink(missing_ok=True)
            raise ValueError("SQLite file has no user tables")
        return dest, tables

    if suffix in CSV_EXTENSIONS:
        table_name, _cols, rows = _read_csv_table(upload, _sanitize_table_name(stem))
        created = _build_sqlite_from_tables({table_name: rows}, dest)
        return dest, created

    if suffix in JSON_EXTENSIONS:
        if suffix in {".jsonl", ".ndjson"}:
            table_name, _cols, rows = _read_jsonl_table(upload, _sanitize_table_name(stem))
            created = _build_sqlite_from_tables({table_name: rows}, dest)
            return dest, created
        parsed = _read_json_tables(upload)
        created = _build_sqlite_from_tables(parsed, dest)
        return dest, created

    if suffix in EXCEL_EXTENSIONS:
        if suffix == ".xls":
            raise ValueError("Legacy .xls is not supported; save as .xlsx")
        parsed = _read_excel_tables(upload)
        created = _build_sqlite_from_tables(parsed, dest)
        return dest, created

    raise ValueError("Unsupported file type")


def sqlite_config_name(path: Path) -> str:
    """Relative path stored in helix.config database.name."""
    base = Path(settings.BASE_DIR)
    try:
        rel = path.resolve().relative_to(base.resolve())
        return str(rel).replace("\\", "/")
    except ValueError:
        return str(path)
